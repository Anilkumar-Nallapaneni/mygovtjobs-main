from pathlib import Path
from datetime import date

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(r"D:\My Projects\mygovtjobs-main")
ENV = ROOT / "backend" / ".env"
OUT_DIR = ROOT / "exports"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / f"jobs-export-{date.today().isoformat()}.xlsx"
STABLE = OUT_DIR / "all-jobs.xlsx"


def load_env(path: Path) -> dict:
    vals = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        vals[k.strip()] = v.strip().strip('"').strip("'")
    return vals


env = load_env(ENV)
db_url = env["DATABASE_URL"]
# SQLAlchemy async URLs -> libpq/psycopg
if db_url.startswith("postgresql+asyncpg://"):
    db_url = "postgresql://" + db_url[len("postgresql+asyncpg://") :]
elif db_url.startswith("postgres+asyncpg://"):
    db_url = "postgresql://" + db_url[len("postgres+asyncpg://") :]

cols = [
    "id",
    "slug",
    "title",
    "dept",
    "status",
    "vacancies",
    "last_date",
    "apply_url",
    "pdf_url",
    "category",
    "state",
    "source",
    "source_url",
    "published_at",
    "published_to_site",
]

sql = """
SELECT
  id::text AS id,
  slug,
  title,
  dept,
  status,
  vacancies,
  last_date::text AS last_date,
  apply_url,
  COALESCE(primary_pdf_url, notification_url, pdf_storage_path) AS pdf_url,
  category,
  array_to_string(state_codes, ',') AS state,
  COALESCE(source_domain, source_type) AS source,
  source_url,
  published_at::text AS published_at,
  published_to_site
FROM jobs
ORDER BY
  CASE status WHEN 'live' THEN 0 WHEN 'draft' THEN 1 ELSE 2 END,
  published_at DESC NULLS LAST,
  title
"""

with psycopg.connect(db_url) as conn:
    with conn.cursor() as cur:
        cur.execute(sql)
        rows = cur.fetchall()
        cur.execute(
            "SELECT status, COUNT(*), COALESCE(SUM(vacancies),0) FROM jobs GROUP BY status"
        )
        by_status = {r[0]: (r[1], r[2]) for r in cur.fetchall()}
        cur.execute("SELECT COUNT(*) FROM jobs")
        total = cur.fetchone()[0]
        cur.execute(
            "SELECT COUNT(*), COALESCE(SUM(vacancies),0) FROM jobs WHERE status = 'live'"
        )
        live_cnt, live_vac = cur.fetchone()

wb = Workbook()
ws_all = wb.active
ws_all.title = "All Jobs"
ws_all.append(cols)
for cell in ws_all[1]:
    cell.font = Font(bold=True)
for row in rows:
    ws_all.append(list(row))

ws_live = wb.create_sheet("Live Jobs")
ws_live.append(cols)
for cell in ws_live[1]:
    cell.font = Font(bold=True)
for row in rows:
    if row[4] == "live":
        ws_live.append(list(row))

ws_sum = wb.create_sheet("Summary")
ws_sum.append(["Metric", "Value"])
ws_sum["A1"].font = Font(bold=True)
ws_sum["B1"].font = Font(bold=True)
ws_sum.append(["Jobs available on website (live)", live_cnt])
ws_sum.append(["Total jobs in DB", total])
for status in ("live", "draft", "expired"):
    cnt, vac = by_status.get(status, (0, 0))
    ws_sum.append([f"Status: {status}", cnt])
    ws_sum.append([f"Vacancies sum ({status})", vac])
ws_sum.append(["Vacancy sum (live only)", live_vac])
ws_sum.append(["Export date", date.today().isoformat()])
ws_sum.append(["Excel path", str(OUT)])

wb.save(OUT)
wb.save(STABLE)

print(f"WROTE={OUT}")
print(f"STABLE={STABLE}")
print(f"LIVE={live_cnt}")
print(f"TOTAL={total}")
print(f"LIVE_VAC={live_vac}")
for s, (c, v) in sorted(by_status.items()):
    print(f"STATUS_{s}={c}|vac={v}")
print(f"ROWS_EXPORTED={len(rows)}")
